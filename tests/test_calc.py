"""Инженерный калькулятор (package calc): every formula against a hand calculation or a
textbook case, the units, the materials, the verdict and the suggestions, the chat."""
from __future__ import annotations

import json
import math

import pytest

import calc
from calc import beams, bending, bolts, sections, shafts, units, welds
from calc.core import EXCESS, FAIL, LOW, OK, CalcError, Result, fmt, status_of

G = 9.80665


def check(result: Result, title: str):
    return next(c for c in result.checks if c.title.startswith(title))


def step(result: Result, title: str) -> float:
    return next(s.value for s in result.steps if s.title.startswith(title))


# ── units, numbers, verdict ─────────────────────────────────────────────────

@pytest.mark.parametrize("text, kind, expected", [
    ("5 тс", units.FORCE, 5 * G * 1000), ("500 кгс", units.FORCE, 500 * G), ("20 кН", units.FORCE, 20000),
    ("20", units.FORCE, 20000), ("3 т", units.FORCE, 3 * G * 1000), ("1,5 м", units.LENGTH, 1500),
    ("12 см", units.LENGTH, 120), ("200", units.LENGTH, 200), ("300 Н·м", units.MOMENT, 300_000),
    ("2 кН·м", units.MOMENT, 2e6), ("50 кгс·м", units.MOMENT, 50 * G * 1000), ("300", units.MOMENT, 300_000),
    ("5 кН/м", units.LINE_LOAD, 5.0), ("1500 кгс/м", units.LINE_LOAD, 1.5 * G), ("1 тс/м", units.LINE_LOAD, G),
    ("235 МПа", units.STRESS, 235), ("24 кгс/мм²", units.STRESS, 24 * G),
])
def test_units(text, kind, expected):
    assert units.parse(text, kind) == pytest.approx(expected)


def test_units_refuse_a_wrong_unit():
    with pytest.raises(CalcError, match="единица"):
        units.parse("5 кН/м", units.FORCE)
    with pytest.raises(CalcError, match="не число"):
        units.parse("много", units.FORCE)
    assert units.kind_of("5 кН/м") == units.LINE_LOAD and units.kind_of("2 кН·м") == units.MOMENT
    assert units.kind_of("20 кгс") == units.FORCE and units.kind_of("20") is None


def test_fmt_and_status():
    assert fmt(63.29) == "63,3" and fmt(0.33849) == "0,338" and fmt(50000) == "50000" and fmt(-1.5) == "−1,5"
    assert status_of(0.9, 1.5) == FAIL and status_of(1.2, 1.5) == LOW and status_of(2, 1.5) == OK
    assert status_of(5, 1.5) == EXCESS


# ── materials ───────────────────────────────────────────────────────────────

def test_strength_by_grade_and_thickness():
    from calc import strength

    s = strength.strength("09Г2С", 8)
    assert (s.yield_, s.ultimate, s.elastic) == (345, 490, 206_000)
    assert strength.strength("09Г2С", 25).yield_ == 305
    assert strength.strength("Ст3", 30).yield_ == 235
    assert strength.strength("Д16Т").elastic == 72_000
    assert strength.strength("40Х").alloyed and not strength.strength("Сталь 45").alloyed
    with pytest.raises(CalcError, match="нет σт"):
        strength.strength("латунь")
    stronger = [x.name for x in strength.stronger(strength.strength("Сталь 45"))]
    assert "40Х" in stronger and "Сталь 20" not in stronger and "АМг6" not in stronger


def test_user_material_with_strength(tmp_path, monkeypatch):
    from cad import materials

    entry = {"name": "Сталь 10ХСНД ГОСТ 19281-2014", "short": "10ХСНД", "kind": "solid", "density": 7.85,
             "synonyms": ["10хснд"], "strength": [[20, 390, 530], [0, 355, 470]], "elastic": 206}
    m = materials.from_json(entry, "user")
    assert m.strength_at(12) == (390, 530) and m.strength_at(40) == (355, 470)
    assert materials.from_json(m.to_json()).strength == m.strength
    with pytest.raises(materials.CatalogError, match="больше σв"):
        materials.from_json({**entry, "strength": [[0, 600, 500]]})


# ── welds ───────────────────────────────────────────────────────────────────

def test_fillet_weld_machine():
    r = welds.calculate({"force": "50 кН", "leg": 6, "length": 200, "thickness": 8, "material": "09Г2С"})
    # τ = N/(0,7·k·(l − 2k)) = 50000/(0,7·6·188)
    assert step(r, "Напряжение от силы") == pytest.approx(50000 / (0.7 * 6 * 188))
    c = check(r, "Прочность шва")
    assert c.limit == pytest.approx(0.6 * 345)
    assert c.margin == pytest.approx(0.6 * 345 / (50000 / (0.7 * 6 * 188)))
    assert r.status == OK


def test_fillet_weld_with_moment_and_dynamics():
    r = welds.calculate({"force": "10 кН", "moment": "1 кН·м", "leg": 5, "length": 150, "count": 2,
                         "full_length": "да", "kd": 2, "electrode": "полуавтомат"})
    tn = 2 * 10_000 / (2 * 0.7 * 5 * 150)
    tm = 2 * 1e6 / (2 * 0.7 * 5 * 150 ** 2 / 6)
    assert check(r, "Прочность шва").actual == pytest.approx(math.hypot(tn, tm))
    assert check(r, "Прочность шва").limit == pytest.approx(0.65 * 255)   # Ст3 up to 10 mm, Св-08Г2С


def test_fillet_weld_sp16():
    # СП 16, п. 14.1.16: N/(βf·kf·lw) ≤ Rwf·γc, N/(βz·kf·lw) ≤ Rwz·γc, Rwz = 0,45·Run
    r = welds.calculate({"force": "300 кН", "leg": 8, "length": 200, "count": 2, "basis": "СП 16", "electrode": "Э50",
                         "material": "09Г2С", "thickness": 10})
    lw = 200 - 16
    assert check(r, "Прочность шва по металлу").actual == pytest.approx(300_000 / (2 * 0.7 * 8 * lw))
    assert check(r, "Прочность шва по металлу").limit == 215
    assert check(r, "Прочность шва по границе").limit == pytest.approx(0.45 * 490)
    assert check(r, "Прочность шва по металлу").required == 1.0


def test_butt_weld_and_notes():
    r = welds.calculate({"joint": "стыковой", "force": "200 кН", "thickness": 10, "length": 150, "kd": 1.5})
    c = check(r, "Прочность стыкового")
    assert c.actual == pytest.approx(300_000 / (10 * 130)) and c.limit == pytest.approx(0.9 * 255)
    assert r.status == FAIL and r.suggestions and "220 мм" in r.suggestions[0]
    sp = welds.calculate({"joint": "стыковой", "force": "200 кН", "thickness": 10, "length": 150,
                          "basis": "СП 16", "inspection": "да"})
    assert check(sp, "Прочность стыкового").limit == pytest.approx(255 / 1.05)
    thin = welds.calculate({"force": "5 кН", "leg": 6, "length": 60, "thickness": 4, "thickness_max": 12})
    notes = " ".join(thin.notes)
    assert "больше 1,2·t" in notes and "меньше 4·k и 40 мм" not in notes
    small = welds.calculate({"force": "5 кН", "leg": 4, "length": 40, "thickness": 4, "thickness_max": 12})
    assert "минимального 6 мм" in " ".join(small.notes) and "меньше 4·k и 40 мм" in " ".join(small.notes)


def test_weld_suggestions_are_calculated():
    r = welds.calculate({"force": "100 кН", "leg": 4, "length": 150, "thickness": 10})
    assert r.status in (FAIL, LOW)
    leg = next(s for s in r.suggestions if s.startswith("катет"))
    k = int(leg.split()[1])
    assert welds.calculate({"force": "100 кН", "leg": k, "length": 150, "thickness": 10}).margin >= 1.5
    assert welds.calculate({"force": "100 кН", "leg": k - 1, "length": 150, "thickness": 10}).margin < 1.5
    assert r.verdict().startswith("Запас ") and "катет" in r.verdict()
    with pytest.raises(CalcError, match="катет"):
        welds.calculate({"force": "10 кН", "length": 100})


# ── bolts ───────────────────────────────────────────────────────────────────

def test_bolts_machine():
    r = bolts.calculate({"diameter": "М16", "count": 4, "shear": "100 кН", "plates": 8, "material": "Ст3"})
    qb = 25_000
    assert check(r, "Срез болта").actual == pytest.approx(qb / (math.pi * 16 ** 2 / 4))
    assert check(r, "Срез болта").limit == pytest.approx(0.6 * 640)
    assert check(r, "Смятие").actual == pytest.approx(qb / (16 * 8)) and check(r, "Смятие").limit == pytest.approx(1.2 * 255)
    t = bolts.calculate({"diameter": "20", "count": 2, "tension": "50 кН", "thread_in_shear": "да"})
    assert check(t, "Разрыв").actual == pytest.approx(25_000 / 245) and check(t, "Разрыв").limit == 660   # 8.8 > M16


def test_bolt_tightening_and_friction():
    r = bolts.calculate({"diameter": "М16", "count": 4, "shear": "10 кН", "plates": 8, "torque": "150 Н·м"})
    f0 = 150_000 / (0.2 * 16)
    assert step(r, "Усилие затяжки") == pytest.approx(f0)
    assert check(r, "Затяжка").actual == pytest.approx(1.3 * f0 / 157)
    assert step(r, "Рекомендуемый момент") == pytest.approx(0.2 * 16 * 0.7 * 640 * 157 / 1.3 / 1000)
    assert step(r, "Запас по сдвигу трением") == pytest.approx(0.15 * f0 * 4 / 10_000)


def test_bolts_sp16_bearing_resistance_matches_the_table():
    # СНиП II-23-81*, табл. 59: Run 360 МПа -> Rbp 475 МПа for bolts of accuracy class B
    r = bolts.calculate({"diameter": "М20", "bolt_class": "5.8", "count": 1, "shear": "30 кН", "plates": 10,
                         "basis": "СП 16", "material": "лист 09Г2С"})
    rbp = step(r, "Расчётное сопротивление смятию")
    assert rbp == pytest.approx((0.6 + 410 * 490 / 206_000) * 490)
    assert (0.6 + 410 * 360 / 206_000) * 360 == pytest.approx(475, abs=2.5)   # the table rounds to 5 MPa
    assert check(r, "Срез болта").limit == pytest.approx(200 * math.pi * 400 / 4 * 0.9)
    with pytest.raises(CalcError, match="12.9"):
        bolts.calculate({"diameter": "М20", "bolt_class": "12.9", "shear": "30 кН", "basis": "СП 16"})


def test_bolt_suggestions():
    r = bolts.calculate({"diameter": "М10", "count": 2, "shear": "60 кН", "plates": 10})
    assert r.status in (FAIL, LOW)
    assert any(s.startswith("М") for s in r.suggestions) and any("болтов" in s for s in r.suggestions)
    with pytest.raises(CalcError, match="М7"):
        bolts.calculate({"diameter": "М7", "shear": "1 кН"})


# ── sections ────────────────────────────────────────────────────────────────

def test_sections_against_the_standards():
    tube = sections.parse("труба 80х40х3")       # ГОСТ 30245-2003: A 6,61 см², Ix 52,25 см⁴, Wx 13,06 см³
    assert tube.area / 100 == pytest.approx(6.61, rel=0.01)
    assert tube.ix / 1e4 == pytest.approx(52.25, rel=0.01) and tube.wx / 1e3 == pytest.approx(13.06, rel=0.01)
    square = sections.parse("труба квадратная 60х3")  # ГОСТ 30245: 60×60×3 — Ix 35,13 см⁴
    assert square.ix / 1e4 == pytest.approx(35.13, rel=0.01)
    pipe = sections.parse("труба 57х3,5")
    assert pipe.ix == pytest.approx(math.pi * (57 ** 4 - 50 ** 4) / 64)
    assert sections.parse("швеллер 16").wx == 93.4e3 and sections.parse("16П").ix == 747e4
    assert sections.parse("двутавр 20").sx == 104e3
    bar = sections.parse("полоса 100х10")
    assert bar.wx == pytest.approx(10 * 100 ** 2 / 6) and bar.sx == pytest.approx(10 * 100 ** 2 / 8)
    rod = sections.parse("круг 40")
    assert rod.wx == pytest.approx(math.pi * 40 ** 3 / 32)
    flat = sections.parse("труба 80х40х3 плашмя")
    assert flat.ix == pytest.approx(tube.iy) and flat.name.endswith("(плашмя)")


def test_bent_channel_by_thin_walled_estimate():
    # a thin-walled estimate with sharp corners: the polygon with the bends must be close to it
    c = sections.parse("гнутый швеллер 100х50х4")
    h, b, t = 96, 48, 4    # the centre line
    ix_thin = t * h ** 3 / 12 + 2 * b * t * (h / 2) ** 2
    assert c.ix == pytest.approx(ix_thin, rel=0.06)
    assert c.area == pytest.approx(4 * (96 + 2 * 48), rel=0.04)


def test_sections_unknown_and_bigger():
    with pytest.raises(CalcError, match="не понял"):
        sections.parse("уголок 50х5")
    with pytest.raises(CalcError, match="номера 17"):
        sections.parse("швеллер 17")
    bigger = sections.bigger(sections.parse("труба 80х40х3"))
    assert bigger and all(s.sizes[-1] >= 3 for s in bigger) and all(s.wx > 13e3 or s.ix > 52e4 for s in bigger)
    assert [s.name for s in sections.bigger(sections.parse("швеллер 16"))][0].startswith("швеллер 18")


# ── beams ───────────────────────────────────────────────────────────────────

EI = 206_000 * 1840e4


def test_beam_solver_textbook_cases():
    L, q, P = 6000.0, 10.0, 5000.0
    s = beams.solve(L, [(0, beams.HINGE), (L, beams.HINGE)], [beams.Load("line", q, 0, L)], EI)
    assert s.moment.max() == pytest.approx(q * L ** 2 / 8, rel=1e-6)
    assert s.deflection.max() == pytest.approx(5 * q * L ** 4 / (384 * EI), rel=1e-6)
    assert abs(s.shear).max() == pytest.approx(q * L / 2, rel=1e-6)
    s = beams.solve(2000, [(0, beams.FIXED)], [beams.Load("force", P, 2000, 2000)], EI)     # a console
    assert s.deflection.max() == pytest.approx(P * 2000 ** 3 / (3 * EI), rel=1e-6)
    assert abs(s.moment).max() == pytest.approx(P * 2000, rel=1e-6)
    s = beams.solve(L, [(0, beams.FIXED), (L, beams.FIXED)], [beams.Load("line", q, 0, L)], EI)
    assert abs(s.moment).max() == pytest.approx(q * L ** 2 / 12, rel=1e-6)
    # two equal spans: the middle support takes 1,25·q·l
    s = beams.solve(L, [(0, beams.HINGE), (L / 2, beams.HINGE), (L, beams.HINGE)], [beams.Load("line", q, 0, L)], EI)
    middle = next(r for p, _, r, _ in s.reactions if p == L / 2)
    assert middle == pytest.approx(1.25 * q * L / 2, rel=1e-6)
    with pytest.raises(CalcError, match="не закреплена"):
        beams.solve(L, [(1000, beams.HINGE)], [beams.Load("force", P, 0, 0)], EI)


def test_beam_inputs():
    assert beams.parse_supports("", 5000) == [(0, beams.HINGE), (5000, beams.HINGE)]
    assert beams.parse_supports("1,5 м; заделка 4500", 6000) == [(1500, beams.HINGE), (4500, beams.FIXED)]
    loads = beams.parse_loads("20 кН в 1500; 5 кН/м; 3 тс/м от 0 до 2 м; 2 кН·м в 1000", 6000)
    assert [(l.kind, l.start, l.end) for l in loads] == [("force", 1500, 1500), ("line", 0, 6000), ("line", 0, 2000),
                                                         ("moment", 1000, 1000)]
    assert loads[2].value == pytest.approx(3 * G)
    with pytest.raises(CalcError, match="где приложена"):
        beams.parse_loads("20 кН", 6000)
    with pytest.raises(CalcError, match="за пределами"):
        beams.parse_supports("7000", 6000)


def test_beam_trailer_spar():
    raw = {"length": "7 м", "supports": "1500; 5500", "loads": "3 тс в 0; 1500 кгс/м от 1500 до 7000",
           "profile": "швеллер 16", "beams": 2, "kd": 1.5, "self_weight": "нет"}
    r = beams.calculate(raw)
    # statics: the two reactions carry all the load of one spar
    total = (3 * G * 1000 + 1.5 * G * 5500) / 2
    assert sum(float(row[2].replace(",", ".").replace("−", "-")) for row in r.tables[0][2]) * 1000 == pytest.approx(total, rel=0.01)
    sigma = check(r, "Прочность при изгибе")
    m = step(r, "Наибольший изгибающий момент") * 1e6
    assert sigma.actual == pytest.approx(1.5 * m / 93.4e3)
    assert check(r, "Прогиб").basis.startswith("2·l")        # the drawbar overhang governs the deflection
    assert r.status == FAIL and r.suggestions[0].startswith("швеллер")
    better = sections.parse(r.suggestions[0].split(" —")[0])
    assert beams._calculate({**beams.read(beams.FIELDS, raw).values}, [], better, None, False).status in (OK, EXCESS)


def test_beam_self_weight_and_lighter_profile():
    r = beams.calculate({"length": 3000, "loads": "1 кН в 1500", "profile": "двутавр 30"})
    assert step(r, "Собственный вес") == pytest.approx(36.5 * G / 1000)
    assert r.status == EXCESS and r.suggestions and r.suggestions[0].startswith("хватит двутавр")


# ── sheet bending ───────────────────────────────────────────────────────────

def test_bending_flat_length_matches_kompas_parts():
    from cad import sheet

    r = bending.calculate({"thickness": 4, "material": "лист Ст3", "radius": 4, "flanges": "40; 100; 40"})
    k = sheet.k_factor(4, 4)
    expected = sheet.build_profile([40, 100, 40], [90, 90], 4, 4, k).flat_length
    assert step(r, "Длина развёртки") == pytest.approx(expected)
    assert expected == pytest.approx(32 + 84 + 32 + 2 * math.pi / 2 * (4 + k * 4))
    assert r.status in (OK, EXCESS)


def test_bending_radius_force_and_press():
    r = bending.calculate({"thickness": 4, "material": "09Г2С", "bend_length": 1000, "press": "100 тс", "along": "да"})
    assert step(r, "Раскрытие матрицы") == 32
    assert step(r, "Внутренний радиус") == pytest.approx(0.16 * 32)
    assert step(r, "Наименьший радиус") == pytest.approx(2.0 * 4)
    assert check(r, "Радиус гиба").status == FAIL and "поперёк проката" in r.suggestions[0]
    assert step(r, "Усилие гибки") == pytest.approx(1.33 * 490 * 16 * 1000 / 32 / 1000)
    across = bending.calculate({"thickness": 4, "material": "09Г2С", "radius": 6})
    assert check(across, "Радиус гиба").status != FAIL
    short = bending.calculate({"thickness": 6, "flanges": "20; 100", "radius": 8})
    assert "короче 0,7·V" in " ".join(short.notes)


# ── shafts ──────────────────────────────────────────────────────────────────

def test_shaft_static_and_fatigue():
    r = shafts.calculate({"diameter": 30, "torque": "300 Н·м", "moment": "200 Н·м"})
    w = math.pi * 30 ** 3 / 32
    sigma, tau = 200_000 / w, 300_000 / (2 * w)
    assert check(r, "Статическая").actual == pytest.approx(math.sqrt(sigma ** 2 + 3 * tau ** 2))
    assert check(r, "Статическая").actual == pytest.approx(math.sqrt(200_000 ** 2 + 0.75 * 300_000 ** 2) / w)
    # keyway, σв 600: Kσ 1,9, Kτ 1,55; ε 0,88/0,77 at Ø30; β 0,9
    s_1 = 0.43 * 600
    ks = 1.9 / (0.88 * 0.9)
    s_sigma = s_1 / (ks * sigma)
    assert step(r, "Запас выносливости по изгибу") == pytest.approx(s_sigma)
    psi_t = 0.5 * (0.02 + 2e-4 * 600)
    s_tau = 0.58 * s_1 / (1.55 / (0.77 * 0.9) * tau / 2 + psi_t * tau / 2)
    assert check(r, "Выносливость").margin == pytest.approx(s_sigma * s_tau / math.hypot(s_sigma, s_tau))


def test_shaft_key_and_suggestions():
    assert shafts.key_for(30) == (8, 7, 4.0) and shafts.key_for(50) == (14, 9, 5.5)
    r = shafts.calculate({"diameter": 30, "torque": "300 Н·м", "raiser": "нет", "key": "да", "key_length": 40})
    assert check(r, "Смятие шпонки").actual == pytest.approx(2 * 300_000 / (30 * 3 * 32))
    assert r.suggestions[0].startswith("шпонка длиной")
    fixed = shafts.calculate({"diameter": 60, "moment": "4 кН·м", "rotating": "нет", "raiser": "нет",
                              "material": "40Х", "kd": 2})
    assert check(fixed, "Статическая").limit == 785
    weak = shafts.calculate({"diameter": 25, "moment": "1 кН·м", "raiser": "нет"})
    assert any(s.startswith("Ø") for s in weak.suggestions)


# ── registry, report, chat ──────────────────────────────────────────────────

def test_registry_and_fields():
    assert list(calc.CALCS) == ["weld", "bolt", "beam", "bend", "shaft"]
    for kind, (title, mod) in calc.CALCS.items():
        names = [f.name for f in mod.FIELDS]
        assert len(names) == len(set(names)), kind
    with pytest.raises(CalcError, match="Не хватает"):
        calc.calculate("weld", {})
    with pytest.raises(CalcError, match="Методика"):
        calc.calculate("weld", {"force": 1, "leg": 4, "length": 100, "basis": "по ощущениям"})


def test_report_xlsx_and_html(tmp_path):
    from openpyxl import load_workbook

    from calc import report

    r = beams.calculate({"length": 6000, "loads": "20 кН в 3000", "profile": "двутавр 20"})
    html = report.to_html(r)
    assert r.verdict() in html.replace("&quot;", '"') and "Реакции опор" in html
    path = report.save_xlsx(r, tmp_path / "балка.xlsx")
    ws = load_workbook(path).active
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert ws["A1"].value == r.title and any(t.startswith("Прочность при изгибе") for t in texts)


class FakeBackend:
    def __init__(self, *answers):
        self.answers = list(answers)
        self.calls = []

    def generate(self, messages, gen=None):
        self.calls.append(messages)
        answer = self.answers.pop(0)
        return answer if isinstance(answer, str) else json.dumps(answer, ensure_ascii=False)


def test_chat_asks_for_what_is_missing_then_counts():
    from calc.chat import Conversation, system_prompt

    backend = FakeBackend({"calc": "weld", "params": {"joint": "угловой", "leg": "6 мм", "force": "5 тс"}},
                          {"calc": "weld", "params": {"joint": "угловой", "leg": "6 мм", "force": "5 тс",
                                                       "length": "200 мм", "material": "09Г2С"}})
    chat = Conversation(backend)
    first = chat.step("угловой шов катет 6, 5 тонн")
    assert first.kind == "ask" and "длина шва" in first.text.lower()
    second = chat.step("длина 200, 09Г2С")
    assert second.kind == "result" and second.result.margin > 1 and second.text.startswith("Сварной шов")
    assert "угловой шов катет 6, 5 тонн. длина 200, 09Г2С" in backend.calls[1][1]["content"]
    assert "calc = \"beam\"" in system_prompt() and "kd" in system_prompt()


def test_chat_total_load_and_cancel():
    from calc.chat import Conversation

    backend = FakeBackend({"calc": "beam", "params": {"length": "6 м", "supports": "1,5 м; 4,5 м",
                                                       "loads": "3 тс/6 м", "profile": "швеллер 16", "beams": "2"}})
    reply = Conversation(backend).step("лонжерон 6 м на осях 1,5 и 4,5, груз 3 т по всей длине, два лонжерона")
    assert reply.kind == "result" and reply.params["loads"].endswith("Н/мм")
    q = units.parse(reply.params["loads"], units.LINE_LOAD)
    assert q == pytest.approx(3 * G * 1000 / 6000)
    chat = Conversation(FakeBackend({"calc": "shaft", "params": {}}))
    assert chat.step("посчитай вал").kind == "ask"
    assert chat.step("отмена").kind == "cancel"
    bad = Conversation(FakeBackend("не знаю"))
    with pytest.raises(CalcError):
        bad.step("что-то")
