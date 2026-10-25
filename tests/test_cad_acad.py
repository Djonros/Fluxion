"""Norm control of AutoCAD drawings (cad/drawing.py, cad/stamp.py, cad/acad_check.py)
on DXF files made here with ezdxf — no AutoCAD needed."""
from __future__ import annotations

import pytest

ezdxf = pytest.importorskip("ezdxf")

from cad import acad, acad_check, report, stamp  # noqa: E402
from cad.drawing import MODEL, read  # noqa: E402


@pytest.fixture(autouse=True)
def _decisions_in_tmp(tmp_path, monkeypatch):
    monkeypatch.setattr(report, "decisions_path", lambda: tmp_path / "report_decisions.json")


def _doc():
    doc = ezdxf.new("R2013", setup=True)
    doc.styles.new("ГОСТ", dxfattribs={"font": "gost.shx"})
    for name in ("Основная", "Тонкая", "Текст", "Рамка", "Размеры", "Лишний"):
        doc.layers.add(name)
    return doc


def _frame(msp, k=1.0, x=0.0, y=0.0, w=420, h=297, layer="Рамка"):
    msp.add_lwpolyline([(x, y), (x + w * k, y), (x + w * k, y + h * k), (x, y + h * k)], close=True,
                       dxfattribs={"layer": layer})
    msp.add_lwpolyline([(x + 20 * k, y + 5 * k), (x + (w - 5) * k, y + 5 * k), (x + (w - 5) * k, y + (h - 5) * k),
                        (x + 20 * k, y + (h - 5) * k)], close=True, dxfattribs={"layer": layer})


def _text(msp, text, x, y, h=3.5, style="ГОСТ", layer="Текст"):
    msp.add_text(text, height=h, dxfattribs={"style": style, "layer": layer}).set_placement((x, y))


def _eskd_stamp(msp, k=1.0, x=0.0, y=0.0, w=420, values=None):
    """Form 1 of ГОСТ 2.104 as texts: the corner of the inner frame is (x0, y0)."""
    v = {"designation": "АБВГ.301234.001", "name": "Кронштейн", "scale": "1:1", "developer": "Иванов",
         "date": "12.10.26", "checker": "Петров", "material": "Ст3 ГОСТ 380-2005"}
    v.update(values or {})
    x0, y0 = x + (w - 5 - 185) * k, y + 5 * k

    def t(text, dx, dy, h=2.5):
        if text:
            _text(msp, text, x0 + dx * k, y0 + dy * k, h * k)

    for label, dy in (("Разраб.", 21), ("Пров.", 16), ("Т.контр.", 11), ("Н.контр.", 6), ("Утв.", 1)):
        t(label, 1, dy)
    t(v["developer"], 18, 21)
    t(v["date"], 56, 21)
    t(v["checker"], 18, 16)
    for label, dx in (("Изм.", 1), ("Лист", 8), ("№ докум.", 18), ("Подп.", 41), ("Дата", 56)):
        t(label, dx, 26)
    t(v["designation"], 90, 45, 5)
    t(v["name"], 75, 26, 5)
    t(v["material"], 75, 6, 3.5)
    t("Лит.", 137, 36)
    t("Масса", 155, 36)
    t("Масштаб", 170, 36)
    t("2,5", 156, 26, 3.5)
    t(v["scale"], 171, 26, 3.5)
    t("Лист", 137, 16)
    t("Листов 1", 160, 16)


def _save(doc, tmp_path, name="чертёж.dxf"):
    path = tmp_path / name
    doc.saveas(path)
    return path


def _problems(result):
    return [r["problem"] for r in result.remarks]


def test_formats_and_scales():
    assert stamp.match_format(420, 297) == ("А3", True, 1.0)
    assert stamp.match_format(210 * 2, 297 * 2) == ("А4", False, 2.0)
    assert stamp.match_format(395, 287) == ("А3", True, 1.0)  # the inner frame
    assert stamp.match_format(400, 300) is None
    assert stamp.parse_scale("М 1 : 2,5") == pytest.approx(0.4)
    assert stamp.scale_text(0.5) == "1:2" and stamp.scale_text(2) == "2:1"
    assert stamp.is_standard_scale(0.4) and not stamp.is_standard_scale(1 / 3)
    assert stamp.field_of_tag("ОБОЗНАЧЕНИЕ") == "designation"
    assert stamp.field_of_tag("Листов") == "sheets"
    assert stamp.field_of_tag("ДАТА_РАЗРАБ") == "developer_date"
    assert stamp.label_of("Н. контр.") == (True, "ncontrol")


def test_clean_text_stamp_has_no_remarks(tmp_path):
    doc = _doc()
    msp = doc.modelspace()
    _frame(msp)
    _eskd_stamp(msp)
    msp.add_line((50, 100), (250, 100), dxfattribs={"layer": "Основная"})
    msp.add_linear_dim(base=(50, 120), p1=(50, 100), p2=(250, 100), dimstyle="Standard",
                       dxfattribs={"layer": "Размеры"}).render()
    msp.add_linear_dim(base=(50, 140), p1=(50, 100), p2=(250, 100), text="%%c<>", dimstyle="Standard",
                       dxfattribs={"layer": "Размеры"}).render()
    msp.add_linear_dim(base=(50, 160), p1=(50, 100), p2=(250, 100), text="200 h14", dimstyle="Standard",
                       dxfattribs={"layer": "Размеры"}).render()
    doc.layers.remove("Лишний")
    result = acad_check.check(read(_save(doc, tmp_path)))
    [sheet] = result.sheets
    assert (sheet.space, sheet.format, sheet.landscape, sheet.k) == (MODEL, "А3", True, 1.0)
    s = sheet.stamp
    assert s.kind == "text" and s.norm == stamp.ESKD
    assert s.value("designation") == "АБВГ.301234.001"
    assert s.value("name") == "Кронштейн"
    assert s.value("scale") == "1:1"
    assert s.value("developer") == "Иванов" and s.value("developer_date") == "12.10.26"
    assert s.value("material").startswith("Ст3")
    assert s.value("mass") == "2,5"
    assert _problems(result) == ["Пустые слои"]  # «Тонкая» of the template
    assert result.remarks[0]["title"] == "Слоёв: 1"


def test_empty_fields_wrong_scale_and_dimensions(tmp_path):
    doc = _doc()
    msp = doc.modelspace()
    _frame(msp, k=2)  # the frame twice its size: the drawing is 1:2
    _eskd_stamp(msp, k=2, values={"name": "", "developer": "", "scale": "1:1"})
    msp.add_line((100, 200), (500, 200))  # layer 0
    msp.add_line((100, 250), (500, 250), dxfattribs={"layer": "Основная", "color": 1})
    dim = msp.add_linear_dim(base=(100, 230), p1=(100, 200), p2=(500, 200), text="350",
                             dimstyle="Standard", dxfattribs={"layer": "Размеры"})
    dim.render()
    msp.add_line((5000, 5000), (5100, 5000), dxfattribs={"layer": "Основная"})  # far outside the frame
    _text(msp, "Примечание", 120, 300, h=6, layer="Текст")  # 3 mm on paper
    result = acad_check.check(read(_save(doc, tmp_path)))
    problems = _problems(result)
    assert "Не заполнено «Наименование»" in problems
    assert "Не заполнено «Разраб.»" in problems
    assert "Масштаб в штампе не совпадает" in problems
    assert "Размер с подменённым значением" in problems
    assert "Объекты на слое 0" in problems
    assert "Цвет или тип линии не «по слою»" in problems
    assert "Объекты за рамкой" in problems
    heights = [r for r in result.remarks if r["problem"] == "Высота шрифта не по ГОСТ 2.304"]
    assert [r["title"] for r in heights] == ["Высота 3 мм"]
    assert "Пустые слои" in problems  # «Лишний»
    # «Внимание» first
    levels = [r["level"] for r in result.remarks]
    assert levels == sorted(levels, key=lambda lv: 0 if lv == report.ATTENTION else 1)


def test_block_stamp_on_a_layout_spds(tmp_path):
    doc = _doc()
    block = doc.blocks.new("ШТАМП_СПДС")
    for tag, y in (("ШИФР", 40), ("НАИМЕНОВАНИЕ", 30), ("СТАДИЯ", 20), ("ЛИСТ", 10), ("РАЗРАБ", 0)):
        block.add_attdef(tag, (0, y), dxfattribs={"height": 3.5, "style": "ГОСТ"})
    block.add_lwpolyline([(0, 0), (185, 0), (185, 55), (0, 55)], close=True)
    layout = doc.layouts.new("Лист1")
    layout.page_setup(size=(420, 297), margins=(0, 0, 0, 0), units="mm")
    _frame(layout)
    ref = layout.add_blockref("ШТАМП_СПДС", (230, 5), dxfattribs={"layer": "Рамка"})
    ref.add_auto_attribs({"ШИФР": "123-2026-КЖ", "НАИМЕНОВАНИЕ": "План плиты", "СТАДИЯ": "Р", "ЛИСТ": "",
                          "РАЗРАБ": "Сидоров"})
    layout.add_viewport(center=(200, 150), size=(300, 200), view_center_point=(0, 0), view_height=200 * 3)
    doc.modelspace().add_line((0, 0), (900, 0), dxfattribs={"layer": "Основная"})
    result = acad_check.check(read(_save(doc, tmp_path)))
    [sheet] = result.sheets
    assert sheet.space == "Лист1" and sheet.stamp.kind == "block"
    assert sheet.stamp.norm == stamp.SPDS and result.norm == stamp.SPDS
    assert sheet.stamp.value("designation") == "123-2026-КЖ"
    assert sheet.scale == pytest.approx(1 / 3)
    problems = _problems(result)
    assert "Не заполнено «Лист»" in problems
    assert "Масштаб не по ГОСТ 2.302" in problems  # 1:3
    assert "Не заполнено «Наименование»" not in problems


def test_no_frame_and_missing_xref(tmp_path):
    doc = _doc()
    doc.modelspace().add_line((0, 0), (10, 10), dxfattribs={"layer": "Основная"})
    doc.add_xref_def("C:/нет/подложка.dwg", "Подложка")
    result = acad_check.check(read(_save(doc, tmp_path)))
    problems = _problems(result)
    assert problems[:2] == ["Рамка не найдена", "Внешняя ссылка не найдена"]


def test_marks_are_kept_per_drawing(tmp_path):
    doc = _doc()
    doc.modelspace().add_line((0, 0), (10, 10))
    path = _save(doc, tmp_path)
    first = acad_check.check(read(path))
    assert first.unchecked
    report.save_decisions(str(path), {r["id"]: {"checked": True} for r in first.remarks})
    again = acad_check.check(read(path), report.load_decisions(str(path)))
    assert not again.unchecked


def test_xlsx_and_texts(tmp_path):
    doc = _doc()
    msp = doc.modelspace()
    _frame(msp)
    _eskd_stamp(msp, values={"name": ""})
    result = acad_check.check(read(_save(doc, tmp_path)))
    target = acad_check.save_xlsx(result, acad_check.default_path(result, tmp_path))
    from openpyxl import load_workbook

    wb = load_workbook(target)
    assert wb.sheetnames == ["Замечания", "Штампы", "Слои"]
    assert any(row[1] == "Не заполнено «Наименование»" for row in wb["Замечания"].iter_rows(values_only=True))
    assert "АБВГ.301234.001" in acad_check.summary_text(result)
    assert "«Наименование»" in acad_check.warnings_text(result)
    assert acad_check.default_path(result, tmp_path).name == "чертёж - нормоконтроль (2).xlsx"


def test_open_drawing_dxf_and_errors(tmp_path):
    doc = _doc()
    path = _save(doc, tmp_path)
    assert acad.open_drawing(path, log=lambda _l: None).file == str(path)
    with pytest.raises(acad.DrawingError):
        acad.open_drawing(tmp_path / "нет.dwg", log=lambda _l: None)
    bad = tmp_path / "x.pdf"
    bad.write_text("x")
    with pytest.raises(acad.DrawingError):
        acad.open_drawing(bad, log=lambda _l: None)


def test_dwg_without_autocad_and_oda(tmp_path, monkeypatch):
    dwg = tmp_path / "a.dwg"
    dwg.write_bytes(b"AC1027")
    monkeypatch.setattr(acad, "oda_converter", lambda: None)
    with pytest.raises(acad.DrawingError, match="ODA File Converter"):
        acad.open_drawing(dwg, log=lambda _l: None, use_autocad=False)


def test_frame_block_with_labels_and_values_over_it(tmp_path):
    """A template: the frame and the stamp labels in one block, the values typed over it as texts."""
    doc = _doc()
    template = doc.blocks.new("Формат А4")
    _frame(template, w=210, h=297)
    for label, dx, dy in (("Разраб.", 1, 21), ("Пров.", 1, 16), ("Масштаб", 170, 36), ("Лист", 137, 16)):
        _text(template, label, 20 + dx, 5 + dy, 2.5)
    msp = doc.modelspace()
    msp.add_blockref("Формат А4", (1000, 0), dxfattribs={"xscale": 4, "yscale": 4, "layer": "Рамка"})
    x0, y0 = 1000 + 20 * 4, 5 * 4

    def value(text, dx, dy, h=3.5):
        _text(msp, text, x0 + dx * 4, y0 + dy * 4, h * 4)

    value("АБВГ.000001.002", 90, 45, 5)
    value("Плита", 75, 26, 5)
    value("Иванов", 18, 21, 2.5)
    value("1:4", 171, 26)
    result = acad_check.check(read(_save(doc, tmp_path)))
    [sheet] = result.sheets
    assert (sheet.format, sheet.landscape, sheet.k) == ("А4", False, 4.0)
    s = sheet.stamp
    assert s.kind == "text" and s.value("designation") == "АБВГ.000001.002" and s.value("name") == "Плита"
    assert s.value("developer") == "Иванов" and s.value("scale") == "1:4"
    problems = _problems(result)
    assert "Масштаб в штампе не совпадает" not in problems
    assert "Не заполнено «Пров.»" in problems


def test_acad_norm_control_is_pro(tmp_path, monkeypatch):
    import os

    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    os.environ.setdefault("FLUXION_DISABLE_WEBENGINE", "1")
    import licensing

    assert acad_check.PRO_FEATURE in licensing.PRO_FEATURES
    monkeypatch.setattr(licensing, "feature_enabled", lambda feature: False)
    opened = []
    monkeypatch.setattr(acad, "open_drawing", lambda path: opened.append(path))
    from cad.__main__ import main

    assert main(["--acad", str(tmp_path / "a.dxf")]) == 2 and not opened
    try:
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QApplication

        import desktop_browser.app as appmod
    except ImportError as exc:
        pytest.skip(f"PySide6 unavailable: {exc}")
    QApplication.instance() or QApplication([])
    told = []
    monkeypatch.setattr(appmod.QMessageBox, "information", lambda parent, title, text: told.append(text))
    window = appmod.FluxionWindow(qsettings=QSettings(str(tmp_path / "ui.ini"), QSettings.IniFormat))
    window.prompt.setText(f"/автокад {tmp_path / 'a.dxf'}")
    window.send_message()
    assert not window.generating and getattr(window, "acad_worker", None) is None and not opened
    assert told and "функция Pro" in told[0]
    window.close()
