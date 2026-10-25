"""Norm control of an AutoCAD drawing: what a norm controller would return it for.

``check`` reads the sheets and title blocks (cad/stamp.py) of a ``Drawing``
(cad/drawing.py) and lists the remarks in the shape of cad/report.py: each one
has to be marked by the designer before the report is made, the marks are kept
per drawing (report.load_decisions / save_decisions).

Checks: frame and title block found and filled (ЕСКД: ГОСТ 2.104, СПДС:
ГОСТ Р 21.101), the scale in the stamp against the frame or the view, the
scales of ГОСТ 2.302, fonts and heights of ГОСТ 2.304, objects on layer 0 or
with colour/linetype not «по слою», visible objects on Defpoints, unused
layers, dimensions with the value typed over, missing external references,
objects outside the frames.
"""
from __future__ import annotations

import re
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path

from . import stamp as st
from .drawing import MODEL, Drawing, Entity
from .report import ATTENTION, CHECK, remark_id

# ГОСТ 2.304: sizes of the font (height of a capital), mm
PRO_FEATURE = "acad_check"  # the norm control of a drawing: Fluxion Pro
FONT_SIZES = (1.8, 2.5, 3.5, 5.0, 7.0, 10.0, 14.0, 20.0, 28.0, 40.0)
# fonts of ГОСТ 2.304 / ISO 3098 as AutoCAD names them (shx files and TrueType families)
GOST_FONT_RE = re.compile(r"gost|гост|iso3098|isocp|isoct|isocteur|isocpeur|isocp2|isoct2|isoct3|isocp3|"
                          r"t-flex|type\s*[ab]|тип\s*[аб]|eskd|ескд|spds|спдс|cs_gost", re.IGNORECASE)
NUMBER_RE = re.compile(r"\d+(?:[.,]\d+)?")
LEVEL = {0: ATTENTION, 1: CHECK}


@dataclass
class Result:
    drawing: Drawing
    sheets: list[st.Sheet]
    remarks: list[dict] = field(default_factory=list)
    norm: str = st.ESKD

    @property
    def unchecked(self) -> list[dict]:
        return [r for r in self.remarks if not r["checked"]]

    @property
    def title(self) -> str:
        return self.drawing.title


class _Remarks:
    def __init__(self, drawing: Drawing, decisions: dict):
        self.out: list[dict] = []
        self.drawing = drawing
        self.decisions = decisions

    def add(self, level: str, problem: str, title: str, where: str, message: str, qty: int = 1) -> None:
        rid = remark_id(problem, title, where)
        n = 2
        while any(r["id"] == rid for r in self.out):
            rid = remark_id(problem, title, where) + f"#{n}"
            n += 1
        mark = self.decisions.get(rid, {})
        self.out.append({"id": rid, "level": level, "problem": problem, "message": message, "kind": "Чертёж",
                         "title": title, "where": where, "qty": qty, "file": self.drawing.file,
                         "checked": bool(mark.get("checked")), "purchased": None})


def _space_title(space: str) -> str:
    return "Модель" if space == MODEL else f"Лист «{space}»"


def _sheet_of(sheets: list[st.Sheet], e: Entity) -> st.Sheet | None:
    for s in sheets:
        if s.space == e.space and e.box and st.inside((e.x, e.y), s.box):
            return s
    return None


def _printed(drawing: Drawing, e: Entity) -> bool:
    layer = drawing.layers.get(e.layer)
    return layer is None or layer.printed


def _stamp_checks(add: _Remarks, sheets: list[st.Sheet]) -> None:
    for s in sheets:
        if s.stamp is None:
            add.add(ATTENTION, "Нет штампа", s.title, _space_title(s.space),
                    "В углу рамки нет основной надписи: ни блока с атрибутами, ни текстов в её ячейках. "
                    "Добавьте штамп по форме 1 ГОСТ 2.104 (СПДС — форма 3 ГОСТ Р 21.101).")
            continue
        stamp = s.stamp
        for key, level in st.REQUIRED[stamp.norm]:
            if key in stamp.known and not stamp.value(key):
                title = st.FIELD_TITLES[key]
                add.add(LEVEL[level], f"Не заполнено «{title}»", s.title, _space_title(s.space),
                        f"В штампе ({stamp.norm}) пустая графа «{title}».")
        _scale_check(add, s)


def _scale_check(add: _Remarks, s: st.Sheet) -> None:
    actual = s.scale
    written = st.parse_scale(s.stamp.value("scale")) if s.stamp else None
    where = _space_title(s.space)
    if actual and not st.is_standard_scale(actual):
        add.add(CHECK, "Масштаб не по ГОСТ 2.302", s.title, where,
                f"Вид в масштабе {st.scale_text(actual)}: такого масштаба нет в ГОСТ 2.302 "
                "(1:2; 1:2,5; 1:4; 1:5; 1:10…; 2:1; 2,5:1; 4:1…).")
    if written and actual and abs(written - actual) > actual * 0.01:
        how = ("рамка начерчена в модели с коэффициентом" if s.space == MODEL else "основной видовой экран листа")
        add.add(ATTENTION, "Масштаб в штампе не совпадает", s.title, where,
                f"В штампе «{s.stamp.value('scale')}», а {how} — {st.scale_text(actual)}. "
                "Исправьте графу «Масштаб» или масштаб вида.")
    if s.stamp and written is None and s.stamp.value("scale"):
        add.add(CHECK, "Масштаб в штампе не распознан", s.title, where,
                f"В графе «Масштаб» записано «{s.stamp.value('scale')}»: ожидается вид 1:2 или 2:1.")
    for v in s.viewport_scales[1:]:
        if not st.is_standard_scale(v):
            add.add(CHECK, "Масштаб не по ГОСТ 2.302", f"{s.title}, вид {st.scale_text(v)}", where,
                    f"Видовой экран в масштабе {st.scale_text(v)}: такого масштаба нет в ГОСТ 2.302.")


def _font_checks(add: _Remarks, drawing: Drawing, sheets: list[st.Sheet]) -> None:
    texts = [e for e in drawing.texts() if e.text.strip() and _printed(drawing, e)]
    by_style: OrderedDict[str, int] = OrderedDict()
    for e in texts:
        by_style[e.style] = by_style.get(e.style, 0) + 1
    for style, n in by_style.items():
        font = drawing.styles.get(style, "")
        if not GOST_FONT_RE.search(font or "") and not GOST_FONT_RE.search(style or ""):
            add.add(CHECK, "Шрифт не по ГОСТ 2.304", f"Стиль «{style}» (шрифт «{font or 'не задан'}»)", "",
                    f"Надписей этим стилем: {n}. Для чертежей — шрифт типа А или Б ГОСТ 2.304 "
                    "(например, GOST type A / gost.shx / ISOCPEUR).", n)
    heights: OrderedDict[tuple[str, float], int] = OrderedDict()
    for e in texts:
        if e.height <= 0:
            continue
        s = _sheet_of(sheets, e)
        k = s.k if s and e.space == MODEL else 1.0
        paper = e.height / k
        if any(abs(paper - h) <= h * 0.03 for h in FONT_SIZES):
            continue
        key = (_space_title(e.space), round(paper, 2))
        heights[key] = heights.get(key, 0) + 1
    for (where, h), n in heights.items():
        add.add(CHECK, "Высота шрифта не по ГОСТ 2.304", f"Высота {st._num(h)} мм", where,
                f"Надписей такой высоты (на бумаге): {n}. Ряд ГОСТ 2.304: 1,8; 2,5; 3,5; 5; 7; 10; 14; 20 мм.", n)


def _layer_checks(add: _Remarks, drawing: Drawing) -> None:
    top = [e for e in drawing.entities if e.kind not in ("ATTRIB", "BLOCKTEXT", "VIEWPORT")]
    for space in drawing.spaces:
        objects = [e for e in top if e.space == space]
        zero = [e for e in objects if e.layer == "0"]
        if zero:
            add.add(CHECK, "Объекты на слое 0", f"Объектов: {len(zero)}", _space_title(space),
                    "Слой 0 — для содержимого блоков. Разнесите объекты по слоям (основные, тонкие, осевые, "
                    "размеры, текст): так проще печать и проверка.", len(zero))
        override = [e for e in objects if not (e.color_bylayer and e.linetype_bylayer)]
        if override:
            add.add(CHECK, "Цвет или тип линии не «по слою»", f"Объектов: {len(override)}", _space_title(space),
                    "Свойства заданы объекту, а не слою: при смене слоя или печати по стилям вид не изменится. "
                    "Поставьте «ПоСлою» (ByLayer).", len(override))
        hidden = [e for e in objects if e.layer.lower() == "defpoints" and e.kind != "DIMENSION"]
        if hidden:
            add.add(ATTENTION, "Объекты на слое Defpoints", f"Объектов: {len(hidden)}", _space_title(space),
                    "Слой Defpoints не печатается: эти объекты пропадут на бумаге и в PDF.", len(hidden))
    unused = [name for name in drawing.layers if name not in drawing.used_layers
              and name not in ("0",) and name.lower() != "defpoints" and "|" not in name]
    if unused:
        listed = ", ".join(unused[:15]) + (f" и ещё {len(unused) - 15}" if len(unused) > 15 else "")
        add.add(CHECK, "Пустые слои", f"Слоёв: {len(unused)}", "",
                f"Слои без объектов: {listed}. Удалите их (_PURGE / ОЧИСТИТЬ), если они не из шаблона.",
                len(unused))


def _dimension_checks(add: _Remarks, drawing: Drawing) -> None:
    by_space: OrderedDict[tuple[str, str], list[str]] = OrderedDict()
    for e in drawing.entities:
        if e.kind != "DIMENSION":
            continue
        text = (e.text or "").strip()
        if not text or "<>" in text:
            continue
        numbers = [float(n.replace(",", ".")) for n in NUMBER_RE.findall(re.sub(r"\\[A-Za-z][^;]*;", "", text))]
        if not numbers:
            by_space.setdefault((e.space, "Размер без числа"), []).append(text)
            continue
        if e.value is None:
            continue
        if not any(abs(n - e.value) <= max(0.01, abs(e.value) * 0.002) for n in numbers):
            by_space.setdefault((e.space, "Размер с подменённым значением"), []).append(
                f"«{text}» при измеренном {st._num(e.value)}")
    for (space, problem), items in by_space.items():
        listed = "; ".join(items[:8]) + (f" и ещё {len(items) - 8}" if len(items) > 8 else "")
        if problem == "Размер без числа":
            add.add(CHECK, problem, f"Размеров: {len(items)}", _space_title(space),
                    f"Текст размера заменён надписью без значения: {listed}.", len(items))
        else:
            add.add(ATTENTION, problem, f"Размеров: {len(items)}", _space_title(space),
                    f"Число на размере не совпадает с измеренным: {listed}. Деталь изготовят не по геометрии: "
                    "уберите подмену (оставьте <>), или исправьте геометрию.", len(items))


def _outside_checks(add: _Remarks, drawing: Drawing, sheets: list[st.Sheet]) -> None:
    for space in drawing.spaces:
        frames = [s for s in sheets if s.space == space]
        if not frames:
            continue
        out = [e for e in drawing.in_space(space)
               if e.box and e.kind not in ("ATTRIB", "BLOCKTEXT", "VIEWPORT") and _printed(drawing, e)
               and not any(st.contains(s.box, e.box, s.k) or st.inside((e.x, e.y), s.box) for s in frames)
               and not any(st.contains(e.box, s.box, s.k) for s in frames)]
        if out:
            add.add(CHECK, "Объекты за рамкой", f"Объектов: {len(out)}", _space_title(space),
                    "Эти объекты вне всех рамок: на листы они не попадут, но могут попасть в печать по границам. "
                    "Удалите их или перенесите на нужный лист.", len(out))


def check(drawing: Drawing, decisions: dict | None = None, norm: str | None = None) -> Result:
    """*norm*: ЕСКД or СПДС; None — by the title block (СПДС when it has «Стадия»)."""
    sheets = st.sheets(drawing, norm)
    add = _Remarks(drawing, decisions or {})
    has_content = [space for space in drawing.spaces if any(e.kind != "VIEWPORT" for e in drawing.in_space(space))]
    if not sheets:
        add.add(ATTENTION, "Рамка не найдена", drawing.title, "",
                "Не найдено ни одной рамки стандартного формата (А4…А0 по ГОСТ 2.301) ни в модели, ни на листах. "
                "Проверить штамп и масштаб нельзя.")
    else:
        for space in has_content:
            if space != MODEL and not any(s.space == space for s in sheets):
                add.add(CHECK, "Лист без рамки", _space_title(space), _space_title(space),
                        "На листе есть объекты, но нет рамки стандартного формата.")
    for x in drawing.xrefs:
        if not x.found:
            add.add(ATTENTION, "Внешняя ссылка не найдена", x.name, "",
                    f"Файл «{x.path or x.name}» не найден: на печати её не будет. Укажите путь заново "
                    "(_XREF / ВНССЫЛКИ).")
    _stamp_checks(add, sheets)
    _font_checks(add, drawing, sheets)
    _dimension_checks(add, drawing)
    _layer_checks(add, drawing)
    _outside_checks(add, drawing, sheets)
    norms = [s.stamp.norm for s in sheets if s.stamp]
    result = Result(drawing, sheets, sorted(add.out, key=lambda r: 0 if r["level"] == ATTENTION else 1),
                    norm or (norms[0] if norms else st.ESKD))
    return result


# ── text and Excel ──────────────────────────────────────────────────────────


def warnings_text(result: Result, limit: int = 12) -> str:
    if not result.remarks:
        return "Замечаний нет."
    counts: OrderedDict[str, int] = OrderedDict()
    for r in result.remarks:
        key = f"{r['level']}: {r['problem'].lower()}"
        counts[key] = counts.get(key, 0) + 1
    lines = [f"Замечаний: {len(result.remarks)}"]
    lines += [f"• {k} — {v}" for k, v in counts.items()]
    lines.append("")
    for r in result.remarks[:limit]:
        where = f" ({r['where']})" if r["where"] else ""
        lines.append(f"{r['level']}: {r['title']}{where} — {r['message']}")
    if len(result.remarks) > limit:
        lines.append(f"…и ещё {len(result.remarks) - limit} — полный список на листе «Замечания».")
    return "\n".join(lines)


def summary_text(result: Result, path: Path | None = None) -> str:
    sheets = len(result.sheets)
    stamps = sum(1 for s in result.sheets if s.stamp)
    attention = sum(1 for r in result.remarks if r["level"] == ATTENTION)
    lines = [f"Нормоконтроль {result.title} ({result.norm}): листов {sheets}, штампов {stamps}, "
             f"замечаний {len(result.remarks)} (из них «{ATTENTION}» — {attention})."]
    for s in result.sheets:
        scale = st.scale_text(s.scale)
        lines.append(f"• {s.title}" + (f", масштаб {scale}" if scale else ""))
    if path:
        lines.append(f"Отчёт: {path}")
    return "\n".join(lines)


def default_path(result: Result, folder: Path | None = None) -> Path:
    from .kompas import default_save_dir

    base = Path(folder) if folder else default_save_dir()
    stem = Path(result.drawing.file).stem or "Чертёж"
    for bad in '<>:"/\\|?*':
        stem = stem.replace(bad, "_")
    path = base / f"{stem} - нормоконтроль.xlsx"
    for n in range(2, 100):  # an earlier report may be open in Excel
        if not path.exists():
            break
        path = base / f"{stem} - нормоконтроль ({n}).xlsx"
    return path


def save_xlsx(result: Result, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, title: str, headers: list[str], widths: list[int]):
        ws.title = title
        ws.append([f"Нормоконтроль: {result.title} ({result.norm})"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([result.drawing.file])
        ws.append(headers)
        for i, cell in enumerate(ws[3]):
            cell.font, cell.fill = bold, head_fill
            ws.column_dimensions[cell.column_letter].width = widths[i]
        ws.freeze_panes = "A4"
        return ws

    ws = sheet(wb.active, "Замечания", ["Уровень", "Замечание", "Что", "Где", "Что будет и что сделать", "Кол.",
                                        "Проверено"], [12, 34, 40, 22, 80, 8, 11])
    for r in result.remarks:
        ws.append([r["level"], r["problem"], r["title"], r["where"], r["message"], r["qty"],
                   "да" if r["checked"] else "нет"])
        if r["level"] == ATTENTION:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="B45309")
        for cell in ws[ws.max_row][1:5]:
            cell.alignment = wrap
    if not result.remarks:
        ws.append(["", "Замечаний нет"])

    ws = sheet(wb.create_sheet(), "Штампы", ["Лист", "Формат", "Масштаб (факт)", "Норма", "Штамп", "Графа",
                                             "Значение"], [30, 14, 14, 9, 24, 18, 60])
    for s in result.sheets:
        stamp = s.stamp
        head = [_space_title(s.space), s.format + (" альб." if s.landscape else ""), st.scale_text(s.scale)]
        if stamp is None:
            ws.append(head + ["", "нет", "", ""])
            continue
        kind = f"блок «{stamp.block}»" if stamp.kind == "block" else "тексты в рамке"
        keys = [k for k in st.FIELD_TITLES if k in stamp.known]
        for i, key in enumerate(keys):
            ws.append((head + [stamp.norm, kind] if i == 0 else ["", "", "", "", ""])
                      + [st.FIELD_TITLES[key], stamp.value(key) or "—"])

    ws = sheet(wb.create_sheet(), "Слои", ["Слой", "Объектов", "Цвет", "Тип линии", "Вкл.", "Заморожен",
                                           "Печать"], [32, 10, 8, 18, 7, 10, 8])
    counts: dict[str, int] = {}
    for e in result.drawing.entities:
        if e.kind not in ("ATTRIB", "BLOCKTEXT"):
            counts[e.layer] = counts.get(e.layer, 0) + 1
    for layer in result.drawing.layers.values():
        ws.append([layer.name, counts.get(layer.name, 0), layer.color, layer.linetype, "да" if layer.on else "нет",
                   "да" if layer.frozen else "нет", "да" if layer.plot else "нет"])
    wb.active = 0
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
