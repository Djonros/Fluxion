"""The report of a calculation: HTML for the window and the PDF, a sheet in Excel.

Both have the same parts: the title and the verdict, the method, the inputs,
the calculation step by step (formula, numbers, value), the checks with their
margins, the suggestions and the notes, the tables (reactions of a beam).
Saving a report is a feature of Fluxion Pro (``calc.PRO_FEATURE``).
"""
from __future__ import annotations

import html
from datetime import date
from pathlib import Path

from .core import EXCESS, FAIL, LOW, OK, STATUS_TEXT, Result, fmt

COLORS = {FAIL: "#C62828", LOW: "#B45309", OK: "#2E7D32", EXCESS: "#1565C0"}
FILLS = {FAIL: "FDE2E1", LOW: "FEF3C7", OK: "DCFCE7", EXCESS: "DBEAFE"}


def _e(text) -> str:
    return html.escape(str(text))


def _limit(check) -> str:
    return "—" if check.margin_value is not None else f"{fmt(check.limit)} {check.unit}".strip()


def _actual(check) -> str:
    return "—" if check.margin_value is not None else f"{fmt(check.actual)} {check.unit}".strip()


def to_html(result: Result, diagram_png: str = "") -> str:
    """A self-contained HTML page (Qt rich text: tables, no CSS grid)."""
    color = COLORS[result.status]
    parts = [f"<h2 style='margin-bottom:2px'>{_e(result.title)}</h2>",
             f"<p style='font-size:13pt;color:{color};margin-top:2px'><b>{_e(result.verdict())}</b></p>",
             f"<p style='color:#555'>Методика: {_e(result.basis)}.</p>"]
    if result.inputs:
        parts.append("<h3>Исходные данные</h3><table cellspacing='0' cellpadding='3'>")
        parts += [f"<tr><td style='color:#555'>{_e(k)}</td><td><b>{_e(v)}</b></td></tr>" for k, v in result.inputs]
        parts.append("</table>")
    parts.append("<h3>Расчёт</h3><table cellspacing='0' cellpadding='3' border='0'>")
    for s in result.steps:
        sub = f" = {_e(s.substitution)}" if s.substitution else ""
        parts.append(f"<tr><td>{_e(s.title)}</td><td style='color:#555'>{_e(s.formula)}{sub}</td>"
                     f"<td align='right'><b>{_e(fmt(s.value))}</b></td><td>{_e(s.unit)}</td></tr>")
    parts.append("</table>")
    parts.append("<h3>Проверки</h3><table cellspacing='0' cellpadding='4' border='1' style='border-collapse:collapse'>"
                 "<tr><th>Проверка</th><th>Факт</th><th>Предел</th><th>Запас</th><th>Нужно</th><th>Вывод</th></tr>")
    for c in result.checks:
        basis = f"<br><span style='color:#777'>{_e(c.basis)}</span>" if c.basis else ""
        parts.append(f"<tr><td>{_e(c.title)}</td><td align='right'>{_e(_actual(c))}</td>"
                     f"<td align='right'>{_e(_limit(c))}{basis}</td><td align='right'><b>{_e(fmt(c.margin, 2))}</b></td>"
                     f"<td align='right'>{_e(fmt(c.required, 2))}</td>"
                     f"<td style='color:{COLORS[c.status]}'><b>{_e(STATUS_TEXT[c.status])}</b></td></tr>")
    parts.append("</table>")
    if diagram_png:
        parts.append(f"<p><img src='{_e(diagram_png)}'></p>")
    for title, header, rows in result.tables:
        parts.append(f"<h3>{_e(title)}</h3><table cellspacing='0' cellpadding='3' border='1' "
                     "style='border-collapse:collapse'><tr>" + "".join(f"<th>{_e(h)}</th>" for h in header) + "</tr>")
        parts += ["<tr>" + "".join(f"<td align='right'>{_e(c)}</td>" for c in row) + "</tr>" for row in rows]
        parts.append("</table>")
    if result.suggestions:
        parts.append("<h3>Что можно сделать</h3><ul>" + "".join(f"<li>{_e(s)}</li>" for s in result.suggestions)
                     + "</ul>")
    if result.notes:
        parts.append("<h3>Примечания</h3><ul>" + "".join(f"<li>{_e(n)}</li>" for n in result.notes) + "</ul>")
    parts.append(f"<p style='color:#777'>Fluxion, инженерный калькулятор, {date.today():%d.%m.%Y}. "
                 "Расчёт по формулам; результат проверяет инженер.</p>")
    return "\n".join(parts)


def save_xlsx(result: Result, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE4EE")
    wb = Workbook()
    ws = wb.active
    ws.title = "Расчёт"
    for col, width in zip("ABCDEF", (46, 44, 16, 12, 10, 18)):
        ws.column_dimensions[col].width = width
    ws.append([result.title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([result.verdict()])
    ws.cell(row=2, column=1).font = Font(bold=True, color=COLORS[result.status].lstrip("#"))
    ws.append([f"Методика: {result.basis}."])
    ws.append([f"Fluxion, инженерный калькулятор, {date.today():%d.%m.%Y}"])

    def header(*titles: str) -> None:
        ws.append([])
        ws.append(list(titles))
        for cell in ws[ws.max_row]:
            cell.font = bold
            cell.fill = head_fill

    if result.inputs:
        header("Исходные данные", "Значение")
        for k, v in result.inputs:
            ws.append([k, v])
    header("Величина", "Формула", "Значение", "Ед.")
    for s in result.steps:
        formula = s.formula + (f" = {s.substitution}" if s.substitution else "")
        ws.append([s.title, formula, round(s.value, 4), s.unit])
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True)
    header("Проверка", "Факт", "Предел", "Запас", "Нужно", "Вывод")
    for c in result.checks:
        limit = _limit(c) + (f" ({c.basis})" if c.basis else "")
        ws.append([c.title, _actual(c), limit, round(c.margin, 2), c.required, STATUS_TEXT[c.status]])
        ws.cell(row=ws.max_row, column=6).fill = PatternFill("solid", fgColor=FILLS[c.status])
    for title, head, rows in result.tables:
        header(title, *head[1:])
        ws.cell(row=ws.max_row, column=1).value = f"{title}: {head[0]}"
        for row in rows:
            ws.append(row)
    if result.suggestions:
        header("Что можно сделать")
        for s in result.suggestions:
            ws.append([s])
    if result.notes:
        header("Примечания")
        for n in result.notes:
            ws.append([n])
    path = Path(path)
    wb.save(path)
    return path


def save_pdf(result: Result, path: Path, diagram_png: str = "") -> Path:
    """The HTML report printed to an A4 PDF by Qt (a QApplication must exist)."""
    from PySide6.QtCore import QMarginsF, QSizeF
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument

    path = Path(path)
    writer = QPdfWriter(str(path))
    writer.setPageLayout(QPageLayout(QPageSize(QPageSize.A4), QPageLayout.Portrait, QMarginsF(15, 15, 15, 15),
                                     QPageLayout.Millimeter))
    writer.setResolution(96)
    doc = QTextDocument()
    doc.setDefaultStyleSheet("body { font-family: 'Segoe UI', Arial; font-size: 9pt; } h2 { font-size: 14pt; }")
    doc.setHtml(f"<body>{to_html(result, diagram_png)}</body>")
    layout = writer.pageLayout()
    rect = layout.paintRectPixels(writer.resolution())
    doc.setPageSize(QSizeF(rect.width(), rect.height()))
    doc.print_(writer)
    return path
